#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Metric rank-consistency matrix -- avg half-matrix version (common metrics / alphabetical order).
Within each dataset: convert each method's value on a metric to a rank, compute pairwise
Spearman rank correlation between metrics, producing a symmetric "metric x metric" matrix,
averaged across datasets and drawn as a half-matrix heatmap.

This version:
  0. [Cleaning unified with the paper] the omics sheet is cleaned on the same basis as the
     main analysis before computing:
       . column-misalignment fix: 20 rows where uses_platform_prior is numeric (not
         True/False); assign_accuracy actually lives in the median_diameter_um column, moved back
       . drop GeneSegNet d40 tuning rows
       . drop ROI-subset rows whose GT scope differs (note contains "ROI subset" /
         "GT scope differs" etc., or GT deviates from the dataset mode by >10%)
       . mean over multiple rows of the same dataset+method
     The computation matches the paper: within each dataset rank the methods, pairwise
     Spearman between metrics, average across datasets; datasets with < 4 methods are excluded.
  1. Each file uses its own full metric set, no intersection
       file1 (image sheet): mask-level metrics in IMAGE_CORE
       file2 (omics sheet): transcript-level metrics in CORE
  2. Metrics are ordered alphabetically (by original metric name)
  3. PDF-only output, named relation_1.pdf and relation_2.pdf
  4. [Adapt] file1 reads the single "Image Segmentation Summary" sheet, groups by the dataset
     column, and lowercases column names
  5. [Adapt] dataset aliases merged (different spellings of the same dataset, see ALIAS)
  6. [Adapt] xenium_liver rows with GT~10,654 are actually CosMx human brain data, reassigned
     to cosmx_human_brain (matches nothing if the local data is already fixed)

Style:
  . lower triangle (incl. diagonal): Blues blocks + very thin white outlines (WHITE_LINE_W=0.2);
    the diagonal is the deepest blue (=1) with no number
  . upper triangle (strict, no diagonal): white background + thin gray full outline + number
  . no figure title; row/column names bold and larger

Metrics (15; error/distance ones are lower-is-better and auto-reversed):
  precision recall f1 count_pearson count_spearman count_mae count_rmse
  vec_cosine vec_pearson vec_js_dist assign_overlap assign_accuracy
  loc_mean_um loc_median_um loc_p95_um
  (metrics not present in the file are dropped automatically)

Dataset name merging: 'xenium_2' and 'xenium_2 breast cancer' are the same dataset;
             'merfish_2 liver1' and 'new_merfish_2 liver1' are merged too (multiple rows of
             the same method are averaged to dedup).

Output (yourpath):
  relation_1.pdf
  relation_2.pdf

Dependencies: pandas numpy scipy matplotlib openpyxl
Run: python plot_consistency_final.py
"""

import os, re
import numpy as np
import numpy.ma as ma
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.font_manager as fm
from scipy.stats import spearmanr
import openpyxl

# ---------------- Paths (yourpath placeholders) ----------------
DESKTOP = "yourpath"
FILES = {
    "file1_eval": os.path.join(DESKTOP, "benchmark_result_image.xlsx"),
    "file2_full": os.path.join(DESKTOP, "benchmark_result_omics.xlsx"),
}
OUTDIR = DESKTOP          # save figures next to the data

# ---------------- Configuration ----------------
CORE = ['precision','recall','f1',
        'count_pearson','count_spearman','count_mae','count_rmse',
        'vec_cosine','vec_pearson','vec_js_dist',
        'assign_overlap','assign_accuracy',
        'loc_mean_um','loc_median_um','loc_p95_um']
# image sheet metrics (column names lowercased on read; all higher-is-better)
IMAGE_CORE = ['dice','iou','pixel accuracy',
              'boundary f1','boundary precision','boundary recall',
              'f1','precision','recall',
              'dq','sq','pq','aji','mda(ap0.5:0.95)']
# lower-is-better -> negate before ranking (error/distance metrics)
SMALLER_BETTER = {'vec_js_dist','count_mae','count_rmse',
                  'loc_mean_um','loc_median_um','loc_p95_um'}
MIN_N   = 4                            # min methods per dataset; below this skip (rank correlation is meaningless)
CMAP    = "Blues"
LINE_W  = 0.5                          # upper-triangle gray outline width
WHITE_LINE_W = 0.2                     # white line width between lower-triangle blocks (thinner)
GRID_GRAY = "#B8B8B8"                  # upper-triangle thin gray outline

PRETTY = {
    'precision':'Precision','recall':'Recall','f1':'F1',
    'count_pearson':'Count PCC','count_spearman':'Count Spearman',
    'count_mae':'Count MAE','count_rmse':'Count RMSE',
    'vec_cosine':'Vec Cosine','vec_pearson':'Vec PCC','vec_js_dist':'Vec JS-dist',
    'assign_overlap':'Assign Overlap','assign_accuracy':'Assign Acc',
    'loc_mean_um':'Loc Mean','loc_median_um':'Loc Median','loc_p95_um':'Loc P95',
    'dice':'Dice','iou':'IoU','pixel accuracy':'Pixel Acc',
    'boundary f1':'Boundary F1','boundary precision':'Boundary Prec','boundary recall':'Boundary Rec',
    'dq':'DQ','sq':'SQ','pq':'PQ','aji':'AJI','mda(ap0.5:0.95)':'mDA',
}

# ---------------- Font ----------------
def pick_font():
    cands=["Microsoft YaHei","Noto Sans CJK SC","SimHei","DejaVu Sans"]
    avail={f.name for f in fm.fontManager.ttflist}
    for c in cands:
        if c in avail: return c
    return "DejaVu Sans"
FONT=pick_font()
plt.rcParams.update({"font.family":FONT,"axes.unicode_minus":False})
print("Font:",FONT)

# ---------------- Dataset name merging ----------------
# Adapt: alias table for different spellings of the same dataset (keys are lowercase originals)
ALIAS = {
    'xenium_pancreas_cancer': 'xenium_1',
    'merfish_uterine':        'merfish_uterian',
    'cosmx_lymph_node':       'cosmx_2',      # Cellpose-3 spelling; the other image methods write cosmx_2_lymph_node
    'merfish_mouse_brain':    'merfish_1',    # Cellpose-3 spelling; GT matches merfish_1
    'merfish_liver1':         'merfish_2',
    'merfish_liver2':         'merfish_3',
}
# Adapt: xenium_liver rows with GT~10,654 are actually CosMx human brain data
LIVER_FIX_GT, LIVER_FIX_TOL, LIVER_FIX_TO = 10654, 50, 'cosmx_human_brain'

def norm_dataset_key(name):
    """
    'new_xenium_4 liver'     -> 'xenium_4'
    'xenium_2 breast cancer' -> 'xenium_2'
    'merfish_2 liver1'       -> 'merfish_2'
    'cosmx_1 pancreas'       -> 'cosmx_1'
    """
    s=str(name).strip()
    s=re.sub(r'^(new|original)[_\-]?','',s)                 # strip the new/original prefix
    s=ALIAS.get(s.lower(), s)                        # Adapt: alias merge
    m=re.search(r'([A-Za-z]+)[_\-]?(\d+)', s)        # grab platform+number
    return f"{m.group(1).lower()}_{m.group(2)}" if m else s.lower()

# ---------------- Reading ----------------
def sheet_to_df(ws):
    rows=list(ws.iter_rows(values_only=True))
    hi=None
    for i,r in enumerate(rows[:6]):
        if r and any(str(c).strip().lower()=='method' for c in r if c is not None):
            hi=i;break
    if hi is None: return None
    hdr=[str(c).strip() if c is not None else "" for c in rows[hi]]
    df=pd.DataFrame(rows[hi+1:], columns=hdr)
    df=df.loc[:,~df.columns.duplicated()]            # dedup columns (some sheets repeat the header)
    if 'method' not in df.columns: return None
    df=df[df['method'].notna()]
    df['method']=df['method'].astype(str).str.strip()
    df=df[(df['method']!='')&(df['method'].str.lower()!='method')]
    return df

def merge_by_key(pairs):
    """Merge by norm_dataset_key; vertically concatenate the dfs under the same key."""
    buckets={}
    for raw,df in pairs:
        buckets.setdefault(norm_dataset_key(raw),[]).append(df)
    out=[(k, pd.concat(dfs, ignore_index=True, sort=False)) for k,dfs in buckets.items()]
    out.sort(key=lambda x:x[0])
    return out

def load_file1(path):
    # Adapt: the image sheet has a single "Image Segmentation Summary" sheet; datasets are
    #         distinguished by the dataset column; columns are capitalized (F1 / Precision /
    #         Recall etc.), lowercased to match CORE
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    df=sheet_to_df(wb['Image Segmentation Summary']); wb.close()
    df.columns=[str(c).strip().lower() for c in df.columns]
    if 'dataset' not in df.columns: return []
    # Adapt: xenium_liver rows with GT~10,654 are reassigned to the CosMx human brain
    raw=df['dataset'].astype(str).str.strip().str.lower()
    gtv=pd.to_numeric(df['gt cells'], errors='coerce')
    fix=(raw=='xenium_liver')&((gtv-LIVER_FIX_GT).abs()<=LIVER_FIX_TOL)
    if fix.any():
        print(f"  [Fixed] xenium_liver rows with GT~{LIVER_FIX_GT:,}: {int(fix.sum())} -> {LIVER_FIX_TO}")
        df.loc[fix,'dataset']=LIVER_FIX_TO
    raw=[(ds, g) for ds,g in df.groupby('dataset')]
    return merge_by_key(raw)

NONCOMP_PAT = r"ROI subset|GT scope differs|ROI 6000px|cell count scope differs"
GT_TOL = 0.10

def clean_omics(df):
    """Same omics-sheet cleaning as the main analysis (see header item 0)."""
    df = df.copy()
    df['dataset'] = df['dataset'].astype(str).str.strip()
    upp = df['uses_platform_prior'].astype(str).str.strip() if 'uses_platform_prior' in df.columns \
        else pd.Series('', index=df.index)
    tune = upp.str.contains('d40', na=False)
    if tune.any():
        print(f"  [Clean] dropped {int(tune.sum())} d40 tuning rows")
    df, upp = df[~tune].copy(), upp[~tune]
    shifted = pd.to_numeric(upp, errors='coerce').notna() & ~upp.isin(['True', 'False'])
    if shifted.any() and 'median_diameter_um' in df.columns:
        df['assign_accuracy'] = to_num(df['assign_accuracy'])
        df.loc[shifted, 'assign_accuracy'] = to_num(df.loc[shifted, 'median_diameter_um'])
        print(f"  [Clean] column-misalignment fix on {int(shifted.sum())} rows (assign_accuracy <- median_diameter_um)")
    gt = to_num(df['gt_cells'])
    mode = gt.groupby(df['dataset']).transform(lambda s: s.mode().iloc[0] if s.notna().any() else np.nan)
    bad = (gt / mode - 1).abs() > GT_TOL
    if 'note' in df.columns:
        bad |= df['note'].astype(str).str.contains(NONCOMP_PAT, regex=True)
    if bad.any():
        print(f"  [Clean] dropped {int(bad.sum())} ROI-subset rows with differing GT scope: "
              f"{df.loc[bad, ['method', 'dataset']].astype(str).agg('/'.join, axis=1).tolist()}")
    return df[~bad]

def load_file2(path):
    df = pd.read_excel(path, sheet_name='Summary')
    df = df[df['method'].notna() & df['dataset'].notna()].copy()
    df['method'] = df['method'].astype(str).str.strip().str.replace(r'\s*[(].*?[)]\s*$', '', regex=True)
    df = clean_omics(df)
    raw=[(ds, g) for ds,g in df.groupby('dataset')]
    return merge_by_key(raw)

# ---------------- Computation ----------------
def to_num(s): return pd.to_numeric(s, errors='coerce')

def method_metric_table(df, metrics):
    cols=[m for m in metrics if m in df.columns]
    sub=df[['method']+cols].copy()
    for c in cols: sub[c]=to_num(sub[c])
    return sub.groupby('method')[cols].mean()   # mean over multiple rows of the same method to dedup

def corr_matrix(sub, metrics):
    present=[m for m in metrics if m in sub.columns]
    d=sub.copy()
    for m in present:
        if m in SMALLER_BETTER: d[m]=-d[m]
    M=len(metrics)          # use the full metric list passed in; missing parts stay NaN
    mat=np.full((M,M),np.nan)
    for i in range(M):
        for j in range(M):
            mi, mj = metrics[i], metrics[j]
            if mi not in present or mj not in present:
                continue
            x=d[mi].values; y=d[mj].values
            ok=~(np.isnan(x)|np.isnan(y))
            if ok.sum()>=MIN_N and np.nanstd(x[ok])>0 and np.nanstd(y[ok])>0:
                r,_=spearmanr(x[ok],y[ok]); mat[i,j]=r
    return pd.DataFrame(mat,index=metrics,columns=metrics)

def average_matrix(dfs, metrics):
    mats=[]; used=[]
    for name,df in dfs:
        sub=method_metric_table(df, metrics)
        if sub.shape[0] < MIN_N: continue
        mats.append(corr_matrix(sub, metrics).values)
        used.append((name, sub.shape[0]))
    if not mats: return None, []
    avg=np.nanmean(np.stack(mats,0),0)
    return pd.DataFrame(avg,index=metrics,columns=metrics), used

def present_metrics(dfs, metrics):
    """Return metrics that have at least one non-null value in the datasets (keeping metrics order)"""
    keep=[]
    for m in metrics:
        for _,df in dfs:
            if m in df.columns and to_num(df[m]).notna().any():
                keep.append(m); break
    return keep

# ---------------- Plotting (half matrix / thin outlines, PDF only) ----------------
def draw_heatmap_half(mat, out_base, subtitle=None):
    """
    Lower triangle (incl. diagonal): Blues blocks + thin white outlines; diagonal is the
    deepest blue (=1), no number written.
    Upper triangle (strict): white background + thin gray outlines + number.
    PDF only (out_base has no extension).
    """
    labels=[PRETTY.get(m,m) for m in mat.index]; M=len(labels)
    data=mat.values.astype(float)

    fig,ax=plt.subplots(figsize=(0.72*M+2.6, 0.72*M+2.2))

    # lower triangle (incl. diagonal) Blues blocks
    lower=data.copy()
    keep_lower=np.tril(np.ones((M,M),bool),k=0)    # include diagonal (k=0)
    lower[~keep_lower]=np.nan
    im=ax.imshow(ma.masked_invalid(lower), cmap=CMAP, vmin=0, vmax=1, aspect='equal')

    # full outlines: lower triangle + diagonal get white thin lines; strict upper triangle white bg + gray thin lines
    for i in range(M):
        for j in range(M):
            if j<=i:  # lower triangle + diagonal: transparent fill + white thin outline
                ax.add_patch(mpatches.Rectangle((j-0.5,i-0.5),1,1,
                    facecolor='none', edgecolor="white", linewidth=WHITE_LINE_W, zorder=2))
            else:     # strict upper triangle: white bg + gray thin outline
                ax.add_patch(mpatches.Rectangle((j-0.5,i-0.5),1,1,
                    facecolor="white", edgecolor=GRID_GRAY, linewidth=LINE_W, zorder=2))

    # write numbers only in the strict upper triangle; none on the diagonal
    for i in range(M):
        for j in range(M):
            if j>i:
                v=data[i,j]
                txt="-" if np.isnan(v) else f"{v:.2f}"
                ax.text(j,i,txt,ha='center',va='center',fontsize=13,
                        color="#0B3D91",fontweight='normal',zorder=3)

    ax.set_xlim(-0.5,M-0.5); ax.set_ylim(M-0.5,-0.5)
    ax.set_xticks(range(M)); ax.set_yticks(range(M))
    ax.set_xticklabels(labels, rotation=45, ha='right', fontsize=12, fontweight='bold')
    ax.set_yticklabels(labels, fontsize=12, fontweight='bold')
    ax.tick_params(length=0)
    for s in ax.spines.values(): s.set_visible(False)

    cb=fig.colorbar(im,ax=ax,fraction=0.046,pad=0.03)
    cb.set_label("Spearman rank corr.",fontsize=10)

    if subtitle:
        ax.text(0.5,-0.14,subtitle,transform=ax.transAxes,ha='center',
                va='top',fontsize=8.5,color="#666",style='italic')
    fig.tight_layout()
    # save PDF only
    fig.savefig(f"{out_base}.pdf", dpi=220, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print("  ->", out_base + ".pdf")

# ---------------- Main flow ----------------
def run(tag, dfs, metrics, out_name=None):
    """
    tag      : print label
    dfs      : merged dataset list (name, df)
    metrics  : fixed metric list (already sorted alphabetically)
    out_name : output PDF base name (no extension); None -> default consistency_{tag}_avg
    """
    print(f"\n########## {tag} ##########")
    print("  Merged datasets:", [k for k,_ in dfs])
    print("  Final metric count:", len(metrics))
    print("  Metric list (alphabetical):", metrics)
    avg,used=average_matrix(dfs, metrics)
    if avg is None:
        print("  Not enough data (method count < %d everywhere)"%MIN_N); return
    print("  Datasets averaged (method count):", used)
    print(f"  {len(used)} datasets included in the average")
    print("  Mean correlation matrix:\n" + avg.round(2).to_string())
    sub=f"averaged over {len(used)} datasets (n>={MIN_N} methods each), Spearman on ranks"
    if out_name is None:
        out_name = f"consistency_{tag}_avg"
    draw_heatmap_half(avg, os.path.join(OUTDIR, out_name), subtitle=sub)

if __name__=="__main__":
    # 1. load both files
    print("Loading file1_eval ...")
    dfs1 = load_file1(FILES["file1_eval"])
    print("Loading file2_full ...")
    dfs2 = load_file2(FILES["file2_full"])

    # 2. take each file's full metric set, sorted alphabetically (no intersection)
    metrics1 = sorted(present_metrics(dfs1, IMAGE_CORE))
    metrics2 = sorted(present_metrics(dfs2, CORE))
    print("\nfile1_eval metrics (alphabetical):", metrics1)
    print("file2_full metrics (alphabetical):", metrics2)

    # 3. run each, producing relation_1.pdf and relation_2.pdf
    run("file1_eval", dfs1, metrics1, out_name="relation_1")
    run("file2_full", dfs2, metrics2, out_name="relation_2")

    print("\nDone. Two PDFs saved to yourpath.")
