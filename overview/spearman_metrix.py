#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
指标排序一致性矩阵 (metric rank-consistency) —— avg 半矩阵版（共同指标 / 字典序）。
每个数据集内：把各分割方法在某指标上的值转成秩，指标两两算 Spearman 秩相关，
得到「指标 × 指标」对称矩阵，跨数据集平均，画成半矩阵热力图。

本版本修改：
  1. 两个文件（file1_eval 与 file2_full）使用完全相同的指标集合（取二者交集）
  2. 指标排列顺序按字典序（指标原始名称字母顺序）统一排列
  3. 输出仅 PDF 格式，文件名分别为 relation_1.pdf 和 relation_2.pdf

样式：
  · 下三角(含对角线)：Blues 色块 + 白色细全框线；对角线为最深蓝(=1)，不写数字
  · 上三角(严格,不含对角)：白底 + 灰色细全框线 + 数字
  · 无图标题；行名/列名加粗加大

指标（15 个，误差/距离类越小越好，自动反向对齐）：
  precision recall f1 count_pearson count_spearman count_mae count_rmse
  vec_cosine vec_pearson vec_js_dist assign_overlap assign_accuracy
  loc_mean_um loc_median_um loc_p95_um
  （file2 缺 assign_overlap 会自动剔除，最终以两文件交集为准）

数据集名归并：'xenium_2' 与 'xenium_2 乳腺癌' 视为同一数据集，
             'merfish_2 肝1' 与 '新_merfish_2 肝1' 也归并（同 method 多行取均值去重）。

输出（桌面）：
  relation_1.pdf
  relation_2.pdf

依赖: pandas numpy scipy matplotlib openpyxl
运行: python plot_consistency_final.py
"""

import os, re
import numpy as np
import numpy.ma as ma
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.font_manager as fm
from scipy.stats import spearmanr
import openpyxl

# ---------------- 路径（Windows 桌面）----------------
DESKTOP = r"C:\Users\阿飞\Desktop"
FILES = {
    "file1_eval": os.path.join(DESKTOP, "benchmark_result_image.xlsx"),
    "file2_full": os.path.join(DESKTOP, "benchmark_full_omics.xlsx"),
}
OUTDIR = DESKTOP          # 图直接存桌面

# ---------------- 配置 ----------------
CORE = ['precision','recall','f1',
        'count_pearson','count_spearman','count_mae','count_rmse',
        'vec_cosine','vec_pearson','vec_js_dist',
        'assign_overlap','assign_accuracy',
        'loc_mean_um','loc_median_um','loc_p95_um']
# 越小越好 -> 取负再算秩（误差/距离类）
SMALLER_BETTER = {'vec_js_dist','count_mae','count_rmse',
                  'loc_mean_um','loc_median_um','loc_p95_um'}
MIN_N   = 4                            # 每数据集最少方法数，低于则跳过（秩相关无意义）
CMAP    = "Blues"
LINE_W  = 0.5                          # 全框线线宽（细）
GRID_GRAY = "#B8B8B8"                  # 上三角灰色细框线

PRETTY = {
    'precision':'Precision','recall':'Recall','f1':'F1',
    'count_pearson':'Count PCC','count_spearman':'Count Spearman',
    'count_mae':'Count MAE','count_rmse':'Count RMSE',
    'vec_cosine':'Vec Cosine','vec_pearson':'Vec PCC','vec_js_dist':'Vec JS-dist',
    'assign_overlap':'Assign Overlap','assign_accuracy':'Assign Acc',
    'loc_mean_um':'Loc Mean','loc_median_um':'Loc Median','loc_p95_um':'Loc P95',
}

# ---------------- 字体（中文）----------------
def pick_font():
    cands=["Microsoft YaHei","Noto Sans CJK SC","SimHei","DejaVu Sans"]
    avail={f.name for f in fm.fontManager.ttflist}
    for c in cands:
        if c in avail: return c
    return "DejaVu Sans"
FONT=pick_font()
plt.rcParams.update({"font.family":FONT,"axes.unicode_minus":False})
print("字体:",FONT)

# ---------------- 数据集名归并 ----------------
def norm_dataset_key(name):
    """
    '新_xenium_4 肝'  -> 'xenium_4'
    'xenium_2 乳腺癌' -> 'xenium_2'
    'merfish_2 肝1'   -> 'merfish_2'
    'cosmx_1 胰腺'    -> 'cosmx_1'
    """
    s=str(name).strip()
    s=re.sub(r'^(新|原)[_\-]?','',s)                 # 去中文前缀
    m=re.search(r'([A-Za-z]+)[_\-]?(\d+)', s)        # 抓 平台+编号
    return f"{m.group(1).lower()}_{m.group(2)}" if m else s.lower()

# ---------------- 读取 ----------------
def sheet_to_df(ws):
    rows=list(ws.iter_rows(values_only=True))
    hi=None
    for i,r in enumerate(rows[:6]):
        if r and any(str(c).strip().lower()=='method' for c in r if c is not None):
            hi=i;break
    if hi is None: return None
    hdr=[str(c).strip() if c is not None else "" for c in rows[hi]]
    df=pd.DataFrame(rows[hi+1:], columns=hdr)
    df=df.loc[:,~df.columns.duplicated()]            # 去重列名（个别 sheet 有重复表头）
    if 'method' not in df.columns: return None
    df=df[df['method'].notna()]
    df['method']=df['method'].astype(str).str.strip()
    df=df[(df['method']!='')&(df['method'].str.lower()!='method')]
    return df

def merge_by_key(pairs):
    """按 norm_dataset_key 归并；同一 key 下多个 df 纵向拼接。"""
    buckets={}
    for raw,df in pairs:
        buckets.setdefault(norm_dataset_key(raw),[]).append(df)
    out=[(k, pd.concat(dfs, ignore_index=True, sort=False)) for k,dfs in buckets.items()]
    out.sort(key=lambda x:x[0])
    return out

def load_file1(path):
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    skip=('原tracker','原始','图像','覆盖','待补','指标')
    raw=[]
    for sn in wb.sheetnames:
        if sn.startswith(skip): continue
        df=sheet_to_df(wb[sn])
        if df is not None and len(df)>0:
            raw.append((sn.strip(), df))
    wb.close()
    return merge_by_key(raw)

def load_file2(path):
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    df=sheet_to_df(wb['汇总']); wb.close()
    if 'dataset' not in df.columns: return []
    raw=[(ds, g) for ds,g in df.groupby('dataset')]
    return merge_by_key(raw)

# ---------------- 计算 ----------------
def to_num(s): return pd.to_numeric(s, errors='coerce')

def method_metric_table(df, metrics):
    cols=[m for m in metrics if m in df.columns]
    sub=df[['method']+cols].copy()
    for c in cols: sub[c]=to_num(sub[c])
    return sub.groupby('method')[cols].mean()   # 同 method 多行取均值去重

def corr_matrix(sub, metrics):
    present=[m for m in metrics if m in sub.columns]
    d=sub.copy()
    for m in present:
        if m in SMALLER_BETTER: d[m]=-d[m]
    M=len(metrics)          # 使用传入的完整指标列表，缺失部分为 NaN
    mat=np.full((M,M),np.nan)
    for i in range(M):
        for j in range(M):
            mi, mj = metrics[i], metrics[j]
            if mi not in present or mj not in present:
                continue
            x=d[mi].values; y=d[mj].values
            ok=~(np.isnan(x)|np.isnan(y))
            if ok.sum()>=MIN_N and np.nanstd(x[ok])>0 and np.nanstd(y[ok])>0:
                r,_=spearmanr(x[ok],y[ok]); mat[i,j]=r
    return pd.DataFrame(mat,index=metrics,columns=metrics)

def average_matrix(dfs, metrics):
    mats=[]; used=[]
    for name,df in dfs:
        sub=method_metric_table(df, metrics)
        if sub.shape[0] < MIN_N: continue
        mats.append(corr_matrix(sub, metrics).values)
        used.append((name, sub.shape[0]))
    if not mats: return None, []
    avg=np.nanmean(np.stack(mats,0),0)
    return pd.DataFrame(avg,index=metrics,columns=metrics), used

def present_metrics(dfs, metrics):
    """返回在数据集中至少有一个非空值的指标列表（保持 metrics 原顺序）"""
    keep=[]
    for m in metrics:
        for _,df in dfs:
            if m in df.columns and to_num(df[m]).notna().any():
                keep.append(m); break
    return keep

# ---------------- 画图（半矩阵 / 细框线，仅输出 PDF）----------------
def draw_heatmap_half(mat, out_base, subtitle=None):
    """
    下三角(含对角)：Blues 色块 + 白色细框线；对角为最深蓝(=1)，不写数字
    上三角(严格)：白底 + 灰色细框线 + 数字
    输出仅 PDF（out_base 不含扩展名）
    """
    labels=[PRETTY.get(m,m) for m in mat.index]; M=len(labels)
    data=mat.values.astype(float)

    fig,ax=plt.subplots(figsize=(0.72*M+2.6, 0.72*M+2.2))

    # 下三角(含对角)Blues 色块
    lower=data.copy()
    keep_lower=np.tril(np.ones((M,M),bool),k=0)    # 含对角(k=0)
    lower[~keep_lower]=np.nan
    im=ax.imshow(ma.masked_invalid(lower), cmap=CMAP, vmin=0, vmax=1, aspect='equal')

    # 全框线：下三角+对角白色细线；上三角(严格)白底灰细线
    for i in range(M):
        for j in range(M):
            if j<=i:  # 下三角+对角：透明填充 + 白色细框线
                ax.add_patch(mpatches.Rectangle((j-0.5,i-0.5),1,1,
                    facecolor='none', edgecolor="white", linewidth=LINE_W, zorder=2))
            else:     # 上三角(严格)：白底 + 灰色细框线
                ax.add_patch(mpatches.Rectangle((j-0.5,i-0.5),1,1,
                    facecolor="white", edgecolor=GRID_GRAY, linewidth=LINE_W, zorder=2))

    # 仅严格上三角写数字；对角线不写
    for i in range(M):
        for j in range(M):
            if j>i:
                v=data[i,j]
                txt="–" if np.isnan(v) else f"{v:.2f}"
                ax.text(j,i,txt,ha='center',va='center',fontsize=13,
                        color="#0B3D91",fontweight='normal',zorder=3)

    ax.set_xlim(-0.5,M-0.5); ax.set_ylim(M-0.5,-0.5)
    ax.set_xticks(range(M)); ax.set_yticks(range(M))
    ax.set_xticklabels(labels, rotation=45, ha='right', fontsize=12, fontweight='bold')
    ax.set_yticklabels(labels, fontsize=12, fontweight='bold')
    ax.tick_params(length=0)
    for s in ax.spines.values(): s.set_visible(False)

    cb=fig.colorbar(im,ax=ax,fraction=0.046,pad=0.03)
    cb.set_label("Spearman rank corr.",fontsize=10)

    if subtitle:
        ax.text(0.5,-0.14,subtitle,transform=ax.transAxes,ha='center',
                va='top',fontsize=8.5,color="#666",style='italic')
    fig.tight_layout()
    # 仅保存 PDF
    fig.savefig(f"{out_base}.pdf", dpi=220, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print("  ->", out_base + ".pdf")

# ---------------- 主流程 ----------------
def run(tag, dfs, metrics, out_name=None):
    """
    tag      : 打印标识
    dfs      : 归并后的数据集列表 (name, df)
    metrics  : 固定使用的指标列表（已按字典序排序）
    out_name : 输出 PDF 基本名（不含扩展名），若为 None 则使用默认 consistency_{tag}_avg
    """
    print(f"\n########## {tag} ##########")
    print("  归并后数据集:", [k for k,_ in dfs])
    print("  最终纳入指标数:", len(metrics))
    print("  指标列表(字典序):", metrics)
    avg,used=average_matrix(dfs, metrics)
    if avg is None:
        print("  无足够数据（方法数均 < %d）"%MIN_N); return
    print("  纳入平均的数据集(方法数):", used)
    sub=f"averaged over {len(used)} datasets (n>={MIN_N} methods each), Spearman on ranks"
    if out_name is None:
        out_name = f"consistency_{tag}_avg"
    draw_heatmap_half(avg, os.path.join(OUTDIR, out_name), subtitle=sub)

if __name__=="__main__":
    # 1. 加载两个文件
    print("正在加载 file1_eval ...")
    dfs1 = load_file1(FILES["file1_eval"])
    print("正在加载 file2_full ...")
    dfs2 = load_file2(FILES["file2_full"])

    # 2. 分别计算各自存在的指标集合
    metrics1 = present_metrics(dfs1, CORE)
    metrics2 = present_metrics(dfs2, CORE)
    print("\nfile1_eval 存在指标:", metrics1)
    print("file2_full 存在指标:", metrics2)

    # 3. 取交集，并按字典序排序（使用原始指标名称排序）
    common_metrics = sorted(set(metrics1) & set(metrics2))
    print("\n共同指标（字典序）:", common_metrics)

    if not common_metrics:
        print("错误：两个文件没有共同指标，无法生成可比的热图。")
    else:
        # 4. 使用相同的指标列表分别运行，输出文件名为 relation_1.pdf 和 relation_2.pdf
        run("file1_eval", dfs1, common_metrics, out_name="relation_1")
        run("file2_full", dfs2, common_metrics, out_name="relation_2")

    print("\n完成。两张 PDF 已存到桌面，指标集合与排列完全一致。")